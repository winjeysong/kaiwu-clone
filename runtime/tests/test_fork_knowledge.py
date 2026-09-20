import hashlib
import importlib.util
import json
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from runtime.plugins import fork_knowledge


class KnowledgeSearchTest(unittest.TestCase):
    def test_skips_unreadable_file_without_weakening_read_checks(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            files = {".gitignore": b".DS_Store\n", "answer.md": "知识库可搜索\n".encode()}
            entries = []
            for name, data in files.items():
                (root / name).write_bytes(data)
                entries.append({
                    "path": name,
                    "sha256": hashlib.sha256(data).hexdigest(),
                    "size_bytes": len(data),
                    "source_type": "folder",
                    "source_path": name,
                })
            index = root / "public-index.json"
            index.write_text(json.dumps({"schema_version": 1, "snapshot_id": "test", "files": entries}))

            with patch.object(fork_knowledge, "KNOWLEDGE_ROOT", root), patch.object(fork_knowledge, "INDEX_PATH", index):
                result = json.loads(fork_knowledge.knowledge_search("可搜索"))
                self.assertEqual(result["matches"][0]["path"], "answer.md")
                self.assertIn("file type is not approved", fork_knowledge.knowledge_read(".gitignore"))
                (root / "answer.md").write_text("内容已篡改")
                self.assertIn("does not match", fork_knowledge.knowledge_search("可搜索"))

    def test_extracts_pdf_office_and_image_text(self):
        if not all(importlib.util.find_spec(name) for name in ("pypdfium2", "docx", "openpyxl", "pptx", "PIL")):
            self.skipTest("document libraries are not installed")

        from docx import Document
        from openpyxl import Workbook
        from PIL import Image, ImageDraw, ImageFont
        from pptx import Presentation

        samples = {}
        word = Document()
        word.add_paragraph("WORDMARKER")
        output = BytesIO()
        word.save(output)
        samples["manual.docx"] = (output.getvalue(), "WORDMARKER")

        excel = Workbook()
        excel.active["A1"] = "EXCELMARKER"
        output = BytesIO()
        excel.save(output)
        samples["table.xlsx"] = (output.getvalue(), "EXCELMARKER")

        slides = Presentation()
        slide = slides.slides.add_slide(slides.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, 1_000_000, 1_000_000).text = "SLIDEMARKER"
        output = BytesIO()
        slides.save(output)
        samples["slides.pptx"] = (output.getvalue(), "SLIDEMARKER")

        image = Image.new("RGB", (800, 160), "white")
        ImageDraw.Draw(image).text((30, 35), "HELLO OCR", fill="black", font=ImageFont.load_default(size=72))
        output = BytesIO()
        image.save(output, format="PNG")
        samples["photo.png"] = (output.getvalue(), "HELLO")
        output = BytesIO()
        image.save(output, format="PDF")
        samples["scan.pdf"] = (output.getvalue(), "HELLO")

        for name, (data, marker) in samples.items():
            with self.subTest(name=name):
                text = fork_knowledge.extract_rich(Path(name).suffix, data)
                self.assertIn(marker, text)

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            data = samples["manual.docx"][0]
            (root / "manual.docx").write_bytes(data)
            index = root / "public-index.json"
            index.write_text(json.dumps({
                "schema_version": 1,
                "snapshot_id": "test",
                "files": [{
                    "path": "manual.docx",
                    "sha256": hashlib.sha256(data).hexdigest(),
                    "size_bytes": len(data),
                    "source_type": "git",
                    "source_path": "manual.docx",
                    "repository": "team",
                    "commit": "a" * 40,
                }],
            }))
            with patch.object(fork_knowledge, "KNOWLEDGE_ROOT", root), patch.object(fork_knowledge, "INDEX_PATH", index):
                result = json.loads(fork_knowledge.knowledge_search("WORDMARKER"))
                self.assertEqual(result["matches"][0]["path"], "manual.docx")
                self.assertIn("#extracted-L", result["matches"][0]["citation"])
                self.assertIn("WORDMARKER", fork_knowledge.knowledge_read("manual.docx"))


if __name__ == "__main__":
    unittest.main()
