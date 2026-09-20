"""Extract searchable text from common document and image formats."""

from functools import lru_cache
from io import BytesIO
from pathlib import Path
from subprocess import run
from tempfile import TemporaryDirectory
from zipfile import ZipFile


RICH_SUFFIXES = {
    ".pdf", ".docx", ".xlsx", ".pptx", ".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp", ".gif"
}
MAX_RICH_FILE_BYTES = 16_000_000
MAX_ARCHIVE_BYTES = 64_000_000
MAX_EXTRACTED_CHARS = 1_000_000
MAX_IMAGE_PIXELS = 20_000_000


def _check_office_archive(data):
    with ZipFile(BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 1_000 or sum(entry.file_size for entry in entries) > MAX_ARCHIVE_BYTES:
            raise ValueError("Office 文件解压后过大")


def _ocr(image):
    if image.width * image.height > MAX_IMAGE_PIXELS:
        raise ValueError("图片像素过多")
    with TemporaryDirectory() as directory:
        path = Path(directory) / "image.png"
        image.convert("RGB").save(path)
        result = run(
            ["tesseract", str(path), "stdout", "-l", "chi_sim+eng"],
            capture_output=True,
            timeout=30,
            check=False,
        )
    if result.returncode:
        raise ValueError("图片文字识别失败")
    return result.stdout.decode("utf-8").strip()


@lru_cache(maxsize=4)
def extract_rich(suffix, data):
    try:
        if suffix == ".pdf":
            import pypdfium2 as pdfium

            with pdfium.PdfDocument(data) as pdf:
                if len(pdf) > 100:
                    raise ValueError("PDF 超过 100 页")
                sections = []
                ocr_pages = 0
                for number in range(1, len(pdf) + 1):
                    page = pdf[number - 1]
                    try:
                        text_page = page.get_textpage()
                        try:
                            text = text_page.get_text_range().strip()
                        finally:
                            text_page.close()
                        if not text:
                            ocr_pages += 1
                            if ocr_pages > 20:
                                raise ValueError("PDF 扫描页超过 20 页")
                            width, height = page.get_size()
                            scale = min(2, (MAX_IMAGE_PIXELS / (width * height)) ** 0.5)
                            bitmap = page.render(scale=scale)
                            try:
                                text = _ocr(bitmap.to_pil())
                            finally:
                                bitmap.close()
                        sections.append(f"[第 {number} 页]\n{text}")
                    finally:
                        page.close()
            output = "\n".join(sections)
        elif suffix == ".docx":
            from docx import Document

            _check_office_archive(data)
            lines = []
            for block in Document(BytesIO(data)).iter_inner_content():
                if hasattr(block, "rows"):
                    lines.extend(" | ".join(cell.text.replace("\n", " ") for cell in row.cells) for row in block.rows)
                elif block.text.strip():
                    lines.append(block.text)
            output = "\n".join(lines)
        elif suffix == ".xlsx":
            from openpyxl import load_workbook
            from openpyxl.utils import get_column_letter

            _check_office_archive(data)
            workbook = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
            try:
                lines = []
                for sheet in workbook:
                    lines.append(f"[工作表: {sheet.title}]")
                    for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        if number > 10_000 or len(row) > 256:
                            raise ValueError("Excel 工作表过大")
                        cells = [f"{get_column_letter(column)}={value}" for column, value in enumerate(row, 1) if value is not None]
                        if cells:
                            lines.append(f"第 {number} 行: " + " | ".join(cells))
                output = "\n".join(lines)
            finally:
                workbook.close()
        elif suffix == ".pptx":
            from pptx import Presentation

            _check_office_archive(data)
            lines = []
            for number, slide in enumerate(Presentation(BytesIO(data)).slides, 1):
                lines.append(f"[第 {number} 张幻灯片]")
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text.strip():
                        lines.append(shape.text)
                    if shape.has_table:
                        lines.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
            output = "\n".join(lines)
        else:
            from PIL import Image

            with Image.open(BytesIO(data)) as image:
                output = "[图片 OCR]\n" + _ocr(image)
        if len(output) > MAX_EXTRACTED_CHARS:
            raise ValueError("提取出的文本过长")
        return output
    except (ValueError, TimeoutError):
        raise
    except Exception as error:
        raise ValueError(f"无法解析 {suffix} 文件") from error
